import csv
from datetime import datetime, timedelta
from decimal import Decimal
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.views.generic import ListView, CreateView, UpdateView, DeleteView
from django.utils.decorators import method_decorator
from django.urls import reverse_lazy
from django.db.models import Q, Sum, Count, Avg
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from accounts.mixins import AdminRequiredMixin, User1RequiredMixin
from accounts.models import User
from .models import Attendance, ManDays, Overtime, UploadBatch, AuditLog
from .forms import AttendanceForm, AttendanceFilterForm, ManDaysForm, ManDaysFilterForm, OvertimeForm, OvertimeFilterForm


def log_audit(action, user, company, record_type, record_id, changes=None):
    """Log an audit trail entry"""
    AuditLog.objects.create(
        action=action,
        user=user,
        company=company,
        record_type=record_type,
        record_id=record_id,
        changes=changes or {}
    )


@login_required
def attendance_dashboard(request):
    """Dashboard with attendance metrics and summaries"""
    company = request.current_tenant
    
    today = timezone.now().date()
    week_ago = today - timedelta(days=7)
    month_ago = today - timedelta(days=30)

    context = {
        'total_attendance': Attendance.objects.filter(company=company).count(),
        'present_today': Attendance.objects.filter(
            company=company,
            attendance_date=today,
            status='PRESENT'
        ).count(),
        'pending_approvals': Attendance.objects.filter(
            company=company,
            supervisor_approved=False
        ).count(),
        'mandays_total': ManDays.objects.filter(company=company).aggregate(Sum('days_allocated'))['days_allocated__sum'] or 0,
        'mandays_utilized': ManDays.objects.filter(company=company).aggregate(Sum('days_utilized'))['days_utilized__sum'] or 0,
        'overtime_hours': Overtime.objects.filter(company=company).aggregate(Sum('hours'))['hours__sum'] or 0,
        'overtime_pending': Overtime.objects.filter(
            company=company,
            supervisor_approved=False
        ).count(),
        'latest_uploads': UploadBatch.objects.filter(company=company)[:5],
        'recent_attendance': Attendance.objects.filter(company=company)[:10],
        'recent_mandays': ManDays.objects.filter(company=company)[:10],
        'recent_overtime': Overtime.objects.filter(company=company)[:10],
    }
    
    return render(request, 'attendance/dashboard.html', context)


@login_required
def attendance_list(request):
    """List and filter attendance records with pagination"""
    company = request.current_tenant
    records = Attendance.objects.filter(company=company)
    
    filter_form = AttendanceFilterForm(request.GET, company=company)
    
    if request.GET:
        date_from = request.GET.get('date_from')
        date_to = request.GET.get('date_to')
        employee_id = request.GET.get('employee')
        status = request.GET.get('status')
        cost_center = request.GET.get('cost_center')
        
        if date_from:
            records = records.filter(attendance_date__gte=date_from)
        if date_to:
            records = records.filter(attendance_date__lte=date_to)
        if employee_id:
            records = records.filter(employee_id=employee_id)
        if status:
            records = records.filter(status=status)
        if cost_center:
            records = records.filter(cost_center__icontains=cost_center)
    
    records = records.order_by('-attendance_date')
    
    page = request.GET.get('page', 1)
    per_page = 25
    total = records.count()
    start = (int(page) - 1) * per_page
    end = start + per_page
    
    paginated_records = records[start:end]
    total_pages = (total + per_page - 1) // per_page
    
    context = {
        'records': paginated_records,
        'filter_form': filter_form,
        'current_page': int(page),
        'total_pages': total_pages,
        'total_records': total,
        'has_previous': int(page) > 1,
        'has_next': int(page) < total_pages,
        'previous_page': int(page) - 1 if int(page) > 1 else None,
        'next_page': int(page) + 1 if int(page) < total_pages else None,
    }
    
    return render(request, 'attendance/attendance_list.html', context)


@login_required
def attendance_create(request):
    """Create new attendance record"""
    company = request.current_tenant
    
    if request.method == 'POST':
        form = AttendanceForm(request.POST)
        if form.is_valid():
            record = form.save(commit=False)
            record.company = company
            record.save()
            log_audit('CREATE', request.user, company, 'Attendance', record.id)
            messages.success(request, 'Attendance record created successfully')
            return redirect('attendance:attendance_list')
    else:
        form = AttendanceForm()
    
    context = {'form': form, 'title': 'Create Attendance Record'}
    return render(request, 'attendance/attendance_form.html', context)


@login_required
def attendance_edit(request, pk):
    """Edit attendance record with optimistic locking"""
    company = request.current_tenant
    record = get_object_or_404(Attendance, pk=pk, company=company)
    
    if request.method == 'POST':
        version = int(request.POST.get('version', record.version))
        if version != record.version:
            messages.error(request, 'This record was modified by another user. Please refresh and try again.')
            return redirect('attendance:attendance_edit', pk=pk)
        
        form = AttendanceForm(request.POST, instance=record)
        if form.is_valid():
            old_data = {
                'status': str(record.status),
                'hours_worked': str(record.hours_worked),
            }
            record = form.save(commit=False)
            record.version += 1
            record.save()
            log_audit('UPDATE', request.user, company, 'Attendance', record.id, old_data)
            messages.success(request, 'Attendance record updated successfully')
            return redirect('attendance:attendance_list')
    else:
        form = AttendanceForm(instance=record)
    
    context = {'form': form, 'record': record, 'title': 'Edit Attendance Record'}
    return render(request, 'attendance/attendance_form.html', context)


@login_required
def attendance_delete(request, pk):
    """Delete attendance record"""
    company = request.current_tenant
    record = get_object_or_404(Attendance, pk=pk, company=company)
    
    if request.method == 'POST':
        record_id = record.id
        log_audit('DELETE', request.user, company, 'Attendance', record_id)
        record.delete()
        messages.success(request, 'Attendance record deleted successfully')
        return redirect('attendance:attendance_list')
    
    context = {'record': record}
    return render(request, 'attendance/attendance_confirm_delete.html', context)


@login_required
def attendance_export(request):
    """Export attendance records as CSV or Excel"""
    company = request.current_tenant
    format_type = request.GET.get('format', 'csv')
    
    records = Attendance.objects.filter(company=company)
    
    if format_type == 'excel':
        wb = Workbook()
        ws = wb.active
        ws.title = 'Attendance'
        
        headers = ['Date', 'Employee', 'Status', 'Cost Center', 'Hours Worked', 'Approved', 'Created']
        ws.append(headers)
        
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        
        for record in records:
            ws.append([
                record.attendance_date.strftime('%Y-%m-%d'),
                record.employee.username,
                record.status,
                record.cost_center,
                str(record.hours_worked),
                'Yes' if record.supervisor_approved else 'No',
                record.created_at.strftime('%Y-%m-%d %H:%M'),
            ])
        
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=attendance.xlsx'
        wb.save(response)
        log_audit('EXPORT', request.user, company, 'Attendance', 0, {'format': 'excel'})
        return response
    else:
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename=attendance.csv'
        
        writer = csv.writer(response)
        writer.writerow(['Date', 'Employee', 'Status', 'Cost Center', 'Hours Worked', 'Approved', 'Created'])
        
        for record in records:
            writer.writerow([
                record.attendance_date.strftime('%Y-%m-%d'),
                record.employee.username,
                record.status,
                record.cost_center,
                str(record.hours_worked),
                'Yes' if record.supervisor_approved else 'No',
                record.created_at.strftime('%Y-%m-%d %H:%M'),
            ])
        
        log_audit('EXPORT', request.user, company, 'Attendance', 0, {'format': 'csv'})
        return response


@login_required
def mandays_list(request):
    """List and filter mandays records"""
    company = request.current_tenant
    records = ManDays.objects.filter(company=company)
    
    filter_form = ManDaysFilterForm(request.GET, company=company)
    
    if request.GET:
        date_from = request.GET.get('date_from')
        date_to = request.GET.get('date_to')
        employee_id = request.GET.get('employee')
        project = request.GET.get('project')
        
        if date_from:
            records = records.filter(mandays_date__gte=date_from)
        if date_to:
            records = records.filter(mandays_date__lte=date_to)
        if employee_id:
            records = records.filter(employee_id=employee_id)
        if project:
            records = records.filter(project__icontains=project)
    
    records = records.order_by('-mandays_date')
    
    page = request.GET.get('page', 1)
    per_page = 25
    total = records.count()
    start = (int(page) - 1) * per_page
    end = start + per_page
    
    paginated_records = records[start:end]
    total_pages = (total + per_page - 1) // per_page
    
    summary = ManDays.objects.filter(company=company).aggregate(
        total_allocated=Sum('days_allocated'),
        total_utilized=Sum('days_utilized'),
    )
    
    context = {
        'records': paginated_records,
        'filter_form': filter_form,
        'current_page': int(page),
        'total_pages': total_pages,
        'total_records': total,
        'has_previous': int(page) > 1,
        'has_next': int(page) < total_pages,
        'previous_page': int(page) - 1 if int(page) > 1 else None,
        'next_page': int(page) + 1 if int(page) < total_pages else None,
        'summary': summary,
    }
    
    return render(request, 'attendance/mandays_list.html', context)


@login_required
def mandays_create(request):
    """Create new mandays record"""
    company = request.current_tenant
    
    if request.method == 'POST':
        form = ManDaysForm(request.POST)
        if form.is_valid():
            record = form.save(commit=False)
            record.company = company
            record.save()
            log_audit('CREATE', request.user, company, 'ManDays', record.id)
            messages.success(request, 'ManDays record created successfully')
            return redirect('attendance:mandays_list')
    else:
        form = ManDaysForm()
    
    context = {'form': form, 'title': 'Create ManDays Record'}
    return render(request, 'attendance/mandays_form.html', context)


@login_required
def mandays_edit(request, pk):
    """Edit mandays record with optimistic locking"""
    company = request.current_tenant
    record = get_object_or_404(ManDays, pk=pk, company=company)
    
    if request.method == 'POST':
        version = int(request.POST.get('version', record.version))
        if version != record.version:
            messages.error(request, 'This record was modified by another user. Please refresh and try again.')
            return redirect('attendance:mandays_edit', pk=pk)
        
        form = ManDaysForm(request.POST, instance=record)
        if form.is_valid():
            old_data = {
                'days_allocated': str(record.days_allocated),
                'days_utilized': str(record.days_utilized),
            }
            record = form.save(commit=False)
            record.version += 1
            record.save()
            log_audit('UPDATE', request.user, company, 'ManDays', record.id, old_data)
            messages.success(request, 'ManDays record updated successfully')
            return redirect('attendance:mandays_list')
    else:
        form = ManDaysForm(instance=record)
    
    context = {'form': form, 'record': record, 'title': 'Edit ManDays Record'}
    return render(request, 'attendance/mandays_form.html', context)


@login_required
def mandays_delete(request, pk):
    """Delete mandays record"""
    company = request.current_tenant
    record = get_object_or_404(ManDays, pk=pk, company=company)
    
    if request.method == 'POST':
        record_id = record.id
        log_audit('DELETE', request.user, company, 'ManDays', record_id)
        record.delete()
        messages.success(request, 'ManDays record deleted successfully')
        return redirect('attendance:mandays_list')
    
    context = {'record': record}
    return render(request, 'attendance/mandays_confirm_delete.html', context)


@login_required
def mandays_export(request):
    """Export mandays records as CSV or Excel"""
    company = request.current_tenant
    format_type = request.GET.get('format', 'csv')
    
    records = ManDays.objects.filter(company=company)
    
    if format_type == 'excel':
        wb = Workbook()
        ws = wb.active
        ws.title = 'ManDays'
        
        headers = ['Date', 'Employee', 'Project', 'Allocated', 'Utilized', 'Approved', 'Created']
        ws.append(headers)
        
        header_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        
        for record in records:
            ws.append([
                record.mandays_date.strftime('%Y-%m-%d'),
                record.employee.username,
                record.project,
                str(record.days_allocated),
                str(record.days_utilized),
                'Yes' if record.supervisor_approved else 'No',
                record.created_at.strftime('%Y-%m-%d %H:%M'),
            ])
        
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=mandays.xlsx'
        wb.save(response)
        log_audit('EXPORT', request.user, company, 'ManDays', 0, {'format': 'excel'})
        return response
    else:
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename=mandays.csv'
        
        writer = csv.writer(response)
        writer.writerow(['Date', 'Employee', 'Project', 'Allocated', 'Utilized', 'Approved', 'Created'])
        
        for record in records:
            writer.writerow([
                record.mandays_date.strftime('%Y-%m-%d'),
                record.employee.username,
                record.project,
                str(record.days_allocated),
                str(record.days_utilized),
                'Yes' if record.supervisor_approved else 'No',
                record.created_at.strftime('%Y-%m-%d %H:%M'),
            ])
        
        log_audit('EXPORT', request.user, company, 'ManDays', 0, {'format': 'csv'})
        return response


@login_required
def overtime_list(request):
    """List and filter overtime records"""
    company = request.current_tenant
    records = Overtime.objects.filter(company=company)
    
    filter_form = OvertimeFilterForm(request.GET, company=company)
    
    if request.GET:
        date_from = request.GET.get('date_from')
        date_to = request.GET.get('date_to')
        employee_id = request.GET.get('employee')
        approved = request.GET.get('approved')
        
        if date_from:
            records = records.filter(overtime_date__gte=date_from)
        if date_to:
            records = records.filter(overtime_date__lte=date_to)
        if employee_id:
            records = records.filter(employee_id=employee_id)
        if approved:
            records = records.filter(supervisor_approved=(approved == 'True'))
    
    records = records.order_by('-overtime_date')
    
    page = request.GET.get('page', 1)
    per_page = 25
    total = records.count()
    start = (int(page) - 1) * per_page
    end = start + per_page
    
    paginated_records = records[start:end]
    total_pages = (total + per_page - 1) // per_page
    
    from django.db.models import Case, When, DecimalField
    summary = Overtime.objects.filter(company=company).aggregate(
        total_hours=Sum('hours'),
        approved_hours=Sum(Case(When(supervisor_approved=True, then='hours'), output_field=DecimalField())),
    )
    
    context = {
        'records': paginated_records,
        'filter_form': filter_form,
        'current_page': int(page),
        'total_pages': total_pages,
        'total_records': total,
        'has_previous': int(page) > 1,
        'has_next': int(page) < total_pages,
        'previous_page': int(page) - 1 if int(page) > 1 else None,
        'next_page': int(page) + 1 if int(page) < total_pages else None,
        'summary': summary,
    }
    
    return render(request, 'attendance/overtime_list.html', context)


@login_required
def overtime_create(request):
    """Create new overtime record"""
    company = request.current_tenant
    
    if request.method == 'POST':
        form = OvertimeForm(request.POST)
        if form.is_valid():
            record = form.save(commit=False)
            record.company = company
            record.save()
            log_audit('CREATE', request.user, company, 'Overtime', record.id)
            messages.success(request, 'Overtime record created successfully')
            return redirect('attendance:overtime_list')
    else:
        form = OvertimeForm()
    
    context = {'form': form, 'title': 'Create Overtime Record'}
    return render(request, 'attendance/overtime_form.html', context)


@login_required
def overtime_edit(request, pk):
    """Edit overtime record with optimistic locking"""
    company = request.current_tenant
    record = get_object_or_404(Overtime, pk=pk, company=company)
    
    if request.method == 'POST':
        version = int(request.POST.get('version', record.version))
        if version != record.version:
            messages.error(request, 'This record was modified by another user. Please refresh and try again.')
            return redirect('attendance:overtime_edit', pk=pk)
        
        form = OvertimeForm(request.POST, instance=record)
        if form.is_valid():
            old_data = {
                'hours': str(record.hours),
                'reason': record.reason,
            }
            record = form.save(commit=False)
            record.version += 1
            record.save()
            log_audit('UPDATE', request.user, company, 'Overtime', record.id, old_data)
            messages.success(request, 'Overtime record updated successfully')
            return redirect('attendance:overtime_list')
    else:
        form = OvertimeForm(instance=record)
    
    context = {'form': form, 'record': record, 'title': 'Edit Overtime Record'}
    return render(request, 'attendance/overtime_form.html', context)


@login_required
def overtime_delete(request, pk):
    """Delete overtime record"""
    company = request.current_tenant
    record = get_object_or_404(Overtime, pk=pk, company=company)
    
    if request.method == 'POST':
        record_id = record.id
        log_audit('DELETE', request.user, company, 'Overtime', record_id)
        record.delete()
        messages.success(request, 'Overtime record deleted successfully')
        return redirect('attendance:overtime_list')
    
    context = {'record': record}
    return render(request, 'attendance/overtime_confirm_delete.html', context)


@login_required
def overtime_export(request):
    """Export overtime records as CSV or Excel"""
    company = request.current_tenant
    format_type = request.GET.get('format', 'csv')
    
    records = Overtime.objects.filter(company=company)
    
    if format_type == 'excel':
        wb = Workbook()
        ws = wb.active
        ws.title = 'Overtime'
        
        headers = ['Date', 'Employee', 'Hours', 'Reason', 'Approved', 'Created']
        ws.append(headers)
        
        header_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
        header_font = Font(bold=True, color='000000')
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        
        for record in records:
            ws.append([
                record.overtime_date.strftime('%Y-%m-%d'),
                record.employee.username,
                str(record.hours),
                record.reason,
                'Yes' if record.supervisor_approved else 'No',
                record.created_at.strftime('%Y-%m-%d %H:%M'),
            ])
        
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=overtime.xlsx'
        wb.save(response)
        log_audit('EXPORT', request.user, company, 'Overtime', 0, {'format': 'excel'})
        return response
    else:
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename=overtime.csv'
        
        writer = csv.writer(response)
        writer.writerow(['Date', 'Employee', 'Hours', 'Reason', 'Approved', 'Created'])
        
        for record in records:
            writer.writerow([
                record.overtime_date.strftime('%Y-%m-%d'),
                record.employee.username,
                str(record.hours),
                record.reason,
                'Yes' if record.supervisor_approved else 'No',
                record.created_at.strftime('%Y-%m-%d %H:%M'),
            ])
        
        log_audit('EXPORT', request.user, company, 'Overtime', 0, {'format': 'csv'})
        return response


@login_required
def attendance_bulk_approve(request):
    """Bulk approve attendance records"""
    if request.method == 'POST':
        record_ids = request.POST.getlist('record_ids')
        company = request.current_tenant
        
        updated = Attendance.objects.filter(
            id__in=record_ids,
            company=company
        ).update(
            supervisor_approved=True,
            approved_by=request.user
        )
        
        for record_id in record_ids:
            log_audit('APPROVE', request.user, company, 'Attendance', int(record_id))
        
        messages.success(request, f'{updated} attendance records approved')
    
    return redirect('attendance:attendance_list')


@login_required
def mandays_by_project(request):
    """Summary view: mandays by project"""
    company = request.current_tenant
    
    by_project = ManDays.objects.filter(company=company).values('project').annotate(
        total_allocated=Sum('days_allocated'),
        total_utilized=Sum('days_utilized'),
        count=Count('id')
    ).order_by('-total_allocated')
    
    context = {
        'summary': by_project,
        'total_allocated': sum(p['total_allocated'] for p in by_project),
        'total_utilized': sum(p['total_utilized'] for p in by_project),
    }
    
    return render(request, 'attendance/mandays_by_project.html', context)


@login_required
def overtime_by_employee(request):
    """Summary view: overtime by employee"""
    from django.db.models import Case, When, DecimalField
    company = request.current_tenant
    
    by_employee = Overtime.objects.filter(company=company).values('employee__username').annotate(
        total_hours=Sum('hours'),
        approved_hours=Sum(Case(When(supervisor_approved=True, then='hours'), output_field=DecimalField())),
        count=Count('id')
    ).order_by('-total_hours')
    
    total_hours = Decimal('0')
    approved_hours = Decimal('0')
    for e in by_employee:
        if e['total_hours']:
            total_hours += e['total_hours']
        if e['approved_hours']:
            approved_hours += e['approved_hours']
    
    context = {
        'summary': by_employee,
        'total_hours': total_hours,
        'approved_hours': approved_hours,
    }
    
    return render(request, 'attendance/overtime_by_employee.html', context)
